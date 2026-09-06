"""Bounded-memory BLAST post-processing and transactional output publication."""

from __future__ import annotations

from collections import defaultdict
from contextlib import suppress
from concurrent.futures import ThreadPoolExecutor
from dataclasses import asdict, replace
from datetime import datetime, timezone
import hashlib
import json
import logging
import os
from pathlib import Path
import shutil
import tempfile
import threading
import time

import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell

from . import __version__
from .assignment import TAX_COLS, assign_query
from .io_utils import fasta_stem, read_fasta_order, split_fasta

RAW_COLUMNS = ["unique ID", "Subject ID", "Accession"] + TAX_COLS + [
    "Similarity", "evalue", "query_coverage", "mismatch", "gapopen", "taxonomy_mapped",
    "species_original", "species_uncertainty", "reference_taxonomy_conflict", "reference_missing_nodes"]
TAX_COLUMNS = ["unique ID"] + TAX_COLS + ["Similarity", "query_coverage", "evalue", "Flag",
    "Ambiguous taxa", "assignment_status", "assigned_rank", "hit_count", "candidate_taxa",
    "hit_limit_reached", "reference_uncertainty", "reference_taxonomy_conflict", "reference_missing_nodes"]
INT_COLUMNS = {"mismatch", "gapopen", "hit_count", "candidate_taxa"}
FLOAT_COLUMNS = {"Similarity", "evalue", "query_coverage"}
BOOL_COLUMNS = {"taxonomy_mapped", "hit_limit_reached"}
BLAST_COLUMNS = ["qseqid", "sseqid", "sacc", "saccver", "pident", "evalue", "qcovs", "qcovhsp", "mismatch", "gapopen"]


def sha256_file(path):
    with open(path, "rb") as handle:
        return hashlib.file_digest(handle, "sha256").hexdigest() if hasattr(hashlib, "file_digest") else _hash_stream(handle)


def _hash_stream(handle):
    digest = hashlib.sha256()
    while block := handle.read(1024 * 1024):
        digest.update(block)
    return digest.hexdigest()


class TableSink:
    """Keep only a chunk in memory; split Excel into sheets before its row limit."""

    excel_rows = 1_048_576

    def __init__(self, path, columns, fmt):
        self.path, self.columns, self.fmt = Path(path), columns, fmt
        self.rows = 0
        self.schema = pa.schema([(c, pa.int64() if c in INT_COLUMNS else pa.float64() if c in FLOAT_COLUMNS else pa.bool_() if c in BOOL_COLUMNS else pa.string()) for c in columns])
        self.writer = pq.ParquetWriter(self.path, self.schema, compression="snappy") if fmt == "parquet" else Workbook(write_only=True)
        self.sheet = None
        self.sheet_rows = 0
        self.closed = False

    def append(self, records):
        if not records:
            return
        if self.fmt == "parquet":
            self.writer.write_table(pa.Table.from_pylist(records, schema=self.schema))
        else:
            for record in records:
                if self.sheet is None or self.sheet_rows >= self.excel_rows:
                    self.sheet = self.writer.create_sheet(f"part_{len(self.writer.worksheets) + 1}")
                    self.sheet.append(self.columns)
                    self.sheet_rows = 1
                values = []
                for column in self.columns:
                    value = record.get(column)
                    if isinstance(value, str):
                        if len(value) > 32767:
                            raise ValueError("Excel cell exceeds 32767 characters; use Parquet to preserve the full value")
                        cell = WriteOnlyCell(self.sheet, value=value)
                        cell.data_type = "s"
                        values.append(cell)
                    else:
                        values.append(value)
                self.sheet.append(values)
                self.sheet_rows += 1
        self.rows += len(records)

    def close(self):
        if self.closed:
            return
        if self.fmt == "parquet":
            self.writer.close()
        else:
            if self.sheet is None:
                self.writer.create_sheet("part_1").append(self.columns)
            self.writer.save(self.path)
            self.writer.close()
        self.closed = True


def native_db_prefix(prefix):
    """BLAST 2.17 Windows LMDB cannot open every Unicode DB path."""
    if os.name != "nt" or str(prefix).isascii():
        return prefix
    import ctypes
    from ctypes import wintypes
    kernel = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel.GetShortPathNameW.argtypes = [wintypes.LPCWSTR, wintypes.LPWSTR, wintypes.DWORD]
    kernel.GetShortPathNameW.restype = wintypes.DWORD
    buffer = ctypes.create_unicode_buffer(32768)
    if not kernel.GetShortPathNameW(str(Path(prefix).parent), buffer, len(buffer)) or not buffer.value.isascii():
        raise ValueError("BLAST/LMDB needs an ASCII database path on this Windows volume; move the database to an ASCII path")
    return str(Path(buffer.value) / Path(prefix).name)


def map_chunk(path, tax_dict, opts):
    from .blast_runner import _prefilter_hits, _resolve_sequence_id, _extract_accession
    raw = pd.read_csv(path, sep="\t", names=BLAST_COLUMNS, dtype=str, na_filter=False)
    raw = _prefilter_hits(raw, opts.min_qcov, opts.max_evalue, opts.min_pident)
    records = []
    for values in raw.itertuples(index=False, name=None):
        row = dict(zip(raw.columns, values))
        key, ranks = _resolve_sequence_id(row, tax_dict)
        mapped = ranks is not None
        rank_values = ranks[:7] if mapped else [""] * 7
        records.append({"unique ID": row["qseqid"], "Subject ID": row["sseqid"],
            "Accession": _extract_accession(row["saccver"] or row["sacc"] or row["sseqid"]),
            **dict(zip(TAX_COLS, rank_values)), "Similarity": float(row["pident"]),
            "evalue": float(row["evalue"]), "query_coverage": float(row["qcov"]),
            "mismatch": int(row["mismatch"]), "gapopen": int(row["gapopen"]),
            "taxonomy_mapped": mapped, "species_original": ranks[7] if mapped and len(ranks) > 7 else rank_values[6],
            "species_uncertainty": ranks[8] if mapped and len(ranks) > 8 else "",
            "reference_taxonomy_conflict": ranks[9] if mapped and len(ranks) > 9 else "",
            "reference_missing_nodes": ranks[10] if mapped and len(ranks) > 10 else ""})
    return records


def _publish(staged, work):
    """Roll back earlier replacements if publication of a later file fails."""
    backups, installed = [], []
    try:
        for source, target in staged:
            target.parent.mkdir(parents=True, exist_ok=True)
            if target.exists():
                backup = work / (target.name + ".previous")
                os.replace(target, backup)
                backups.append((backup, target))
            os.replace(source, target)
            installed.append(target)
    except BaseException:
        for target in installed:
            target.unlink(missing_ok=True)
        for backup, target in reversed(backups):
            os.replace(backup, target)
        raise


def run(query_fasta, out_dir, db, options):
    from . import blast_runner as br
    from .dbs import ensure_db_prefix, validate_database
    from .taxmap import find_taxmaps_paths
    from .blast_tools import get_tool_version
    opts = replace(options)
    query = Path(query_fasta).resolve()
    total_queries = len(read_fasta_order(str(query)))
    base = fasta_stem(str(query))
    root = Path(opts.output_dir).resolve() if opts.output_dir else query.parent.parent
    root.mkdir(parents=True, exist_ok=True)
    ext = ".parquet.snappy" if opts.output_format == "parquet" else ".xlsx"
    raw_path = root / "raw_blast" / f"{base}_raw_blast{ext}"
    tax_path = root / "taxonomy" / f"{base}_taxonomy{ext}"
    manifest_path = root / "taxonomy" / f"{base}.runinfo.json"
    legacy_info = root / "taxonomy" / f"{base}.runinfo.txt"
    def existing_outputs():
        return [p for directory, prefix in [(root / "raw_blast", f"{base}_raw_blast."), (root / "taxonomy", f"{base}_taxonomy.")]
                if directory.is_dir() for p in directory.iterdir() if p.name.startswith(prefix)]
    # Output staging is on the same filesystem, so os.replace remains atomic per file.
    lock_path = root / f".{base}.apscale_blast2.lock"
    with lock_path.open("x", encoding="ascii") as handle:
        handle.write(str(os.getpid()))
    work = None
    sinks = []
    finished = False
    started = time.monotonic()
    cancel = threading.Event()
    executor = None
    try:
        collisions = existing_outputs()
        if (collisions or manifest_path.exists() or legacy_info.exists()) and not opts.overwrite:
            raise FileExistsError(f"Outputs already exist for {base}; use --overwrite or a different --out-dir")
        validate_database(db.path)
        prefix = ensure_db_prefix(db.path)
        effective_prefix = native_db_prefix(prefix)
        work = Path(tempfile.mkdtemp(prefix=f".{base}.work_", dir=root))
        subsets = split_fasta(str(query), str(work / "subsets"), opts.subset_size)
        workers = min(opts.workers, len(subsets))
        threads = opts.threads or max(1, (os.cpu_count() or 1) - 2)
        workers = min(workers, threads)
        threads_per = max(1, threads // workers)
        print(f"BLAST: {total_queries} queries, {len(subsets)} chunks, {workers} workers, {threads_per} threads/worker", flush=True)

        def search(fasta):
            return br._run_single(opts.blastn_exe, effective_prefix, fasta, threads_per, opts.task,
                opts.max_target_seqs, opts.masking, opts.log_level == "DEBUG",
                opts.min_pident if opts.inline_perc_identity else None, opts.max_evalue, opts.min_qcov,
                cancel_event=cancel)

        # Check the real database with the first chunk before loading taxonomy or queuing jobs.
        first = search(subsets[0])
        logger = logging.getLogger("apscale_blast2")
        tax_dict = br.get_tax_dict_cached(prefix, logger)
        sinks = [TableSink(work / raw_path.name, RAW_COLUMNS, opts.output_format), TableSink(work / tax_path.name, TAX_COLUMNS, opts.output_format)]
        counts = {"queries": 0, "hits": 0, "unmapped_hits": 0, "queries_at_hit_limit": 0}
        executor = ThreadPoolExecutor(max_workers=workers)
        pending = {}
        next_submit = 1
        for index, subset in enumerate(subsets):
            while next_submit < len(subsets) and len(pending) < workers:
                pending[next_submit] = executor.submit(search, subsets[next_submit])
                next_submit += 1
            tsv = first if index == 0 else pending.pop(index).result()
            records = map_chunk(tsv, tax_dict, opts)
            grouped = defaultdict(list)
            for row in records:
                grouped[row["unique ID"]].append(row)
            output = [assign_query(qid, grouped.get(qid, []), opts) for qid in read_fasta_order(subset)]
            counts["queries"] += len(output)
            counts["hits"] += len(records)
            counts["unmapped_hits"] += sum(not r["taxonomy_mapped"] for r in records)
            counts["queries_at_hit_limit"] += sum(r["hit_limit_reached"] for r in output)
            sinks[0].append(records)
            sinks[1].append(output)
            if not opts.keep_tsv:
                Path(tsv).unlink()
                Path(subset).unlink()
            if index == 0 or (index + 1) % 10 == 0 or index + 1 == len(subsets):
                print(f"Processed {counts['queries']}/{total_queries} queries; {counts['hits']} hits", flush=True)
        if counts["hits"] and counts["unmapped_hits"] == counts["hits"]:
            raise ValueError("None of the BLAST hits map to the taxonomy table; check reference IDs")
        if counts["unmapped_hits"]:
            logger.warning("%s/%s hits have no taxonomy mapping; see assignment_status", counts["unmapped_hits"], counts["hits"])
        if counts["queries_at_hit_limit"]:
            logger.warning("%s queries reached max_target_seqs; additional tied taxa may be unreported", counts["queries_at_hit_limit"])
        for sink in sinks:
            sink.close()
        sinks = []
        effective_options = asdict(opts)
        effective_options.update(threads=threads, workers=workers, threads_per_worker=threads_per)
        metadata = {"version": __version__, "created_at": datetime.now(timezone.utc).isoformat(), "options": effective_options,
            "blastn_version": get_tool_version(opts.blastn_exe),
            "query": str(query), "query_sha256": sha256_file(query), "database_prefix": prefix,
            "native_database_prefix": effective_prefix, "taxonomy_tables": {p: sha256_file(p) for p in find_taxmaps_paths(prefix)},
            "counts": counts, "elapsed_seconds": time.monotonic() - started, "work_dir": str(work) if opts.keep_tsv else None,
            "outputs": {"raw_output": str(raw_path), "taxonomy_output": str(tax_path)}}
        metadata["database_index_metadata"] = [{"name": p.name, "bytes": p.stat().st_size, "mtime_ns": p.stat().st_mtime_ns}
                                               for p in Path(prefix).parent.iterdir() if p.is_file() and p.name.startswith(Path(prefix).name + ".")]
        metadata["implementation_sha256"] = {p.name: sha256_file(p) for p in Path(__file__).parent.glob("*.py")}
        staged_info = work / manifest_path.name
        staged_info.write_text(json.dumps(metadata, indent=2) + "\n", encoding="utf-8")
        staged_text = work / legacy_info.name
        staged_text.write_text("\n".join(f"{k}\t{v}" for k, v in effective_options.items()) + "\n", encoding="utf-8")
        _publish([(work / raw_path.name, raw_path), (work / tax_path.name, tax_path), (staged_text, legacy_info), (staged_info, manifest_path)], work)
        # Explicit overwrite also removes alternate-format tables from this same input.
        for previous in collisions:
            if previous not in {raw_path, tax_path}:
                previous.unlink()
        finished = True
        print(f"Finished in {time.monotonic() - started:.1f}s: {tax_path}", flush=True)
        return {"raw_output": str(raw_path), "taxonomy_output": str(tax_path), "raw_xlsx": str(raw_path), "filtered_xlsx": str(tax_path)}
    finally:
        cancel.set()
        if executor is not None:
            executor.shutdown(wait=True, cancel_futures=True)
        for sink in sinks:
            with suppress(Exception):
                sink.close()
        lock_path.unlink(missing_ok=True)
        if work and finished and not opts.keep_tsv:
            if work.resolve().parent != root.resolve():
                raise RuntimeError("Refusing cleanup outside the output root")
            shutil.rmtree(work)
        elif work and not finished:
            print(f"Run failed; diagnostics retained in {work}", flush=True)
