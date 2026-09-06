"""Inspect official download formats without dumping entire databases."""

import gzip
from pathlib import Path
import tarfile
import zipfile
import openpyxl
import sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

root = Path(__file__).resolve().parents[1] / "dbs" / "official"
for path in root.iterdir():
    if path.name.endswith(".json") or path.name.endswith(".partial"):
        continue
    print("\nFILE", path.name)
    if path.suffix == ".gz":
        with gzip.open(path, "rt", encoding="utf-8") as handle:
            for _ in range(4):
                print(handle.readline().strip()[:2000])
    elif path.suffix == ".tgz":
        with tarfile.open(path) as archive:
            print([m.name for m in archive.getmembers()])
            for member in archive.getmembers():
                if member.isfile() and member.name.endswith((".fasta", ".fa")):
                    with archive.extractfile(member) as handle:
                        print(handle.readline().decode().strip()[:2000])
    elif path.suffix == ".zip":
        with zipfile.ZipFile(path) as archive:
            print(archive.namelist())
            for name in archive.namelist():
                if name.endswith((".fasta", ".fa", ".txt", ".tsv", ".csv")):
                    with archive.open(name) as handle:
                        print(name, handle.readline().decode().strip()[:2000])
    elif path.suffix == ".xlsx":
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        print(book.sheetnames)
        for sheet in book:
            print(sheet.title, sheet.max_row, sheet.max_column)
            print(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))
        book.close()
